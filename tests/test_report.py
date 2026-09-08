"""Tests del resumen, la localización y la exportación a Excel."""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from conftest import bank_frame, bank_row, ledger_frame, ledger_row
from reconciliation.i18n import Translator
from reconciliation.loader import DEFAULT_ACCOUNT, REJECTED_COLUMNS
from reconciliation.matching import reconcile
from reconciliation.report import (
    SHEET_ORDER,
    export_report,
    format_report,
    localize,
    summarize,
    total_flagged,
)


def build(bank_rows, ledger_rows, rejected=None):
    bank = bank_frame(bank_rows)
    ledger = ledger_frame(ledger_rows)
    results = reconcile(bank, ledger)
    summary = summarize(results, bank, ledger, rejected)
    return results, summary


# --------------------------------------------------------------------------
# Totales y control aritmético
# --------------------------------------------------------------------------

def test_totales_por_cuenta():
    _, summary = build(
        [bank_row("2026-08-01", "DEP", "1450.00"), bank_row("2026-08-02", "FEE", "-50.00")],
        [ledger_row("2026-08-01", "Deposito", "1450.00", "INV-1")],
    )
    fila = summary.iloc[0]

    assert fila["bank_total"] == 1400.00
    assert fila["ledger_total"] == 1450.00
    assert fila["difference"] == -50.00
    assert fila["matched"] == 1
    assert fila["only_in_bank"] == 1


def test_la_diferencia_queda_completamente_explicada():
    """`unexplained` es el control que mira primero un contador."""
    _, summary = build(
        [
            bank_row("2026-08-01", "DEP", "1450.00"),
            bank_row("2026-08-06", "GLOBEX", "1800.00"),
            bank_row("2026-08-14", "FEE", "-15.00"),
        ],
        [
            ledger_row("2026-08-01", "Deposito", "1450.00", "INV-1"),
            ledger_row("2026-08-06", "Globex", "1750.00", "INV-2"),
            ledger_row("2026-08-24", "Wayne", "900.00", "INV-3"),
        ],
    )
    fila = summary.iloc[0]

    assert fila["by_discrepancies"] == 50.00
    assert fila["by_only_bank"] == -15.00
    assert fila["by_only_ledger"] == -900.00
    assert fila["unexplained"] == 0.00
    assert fila["difference"] == pytest.approx(
        fila["by_discrepancies"] + fila["by_only_bank"]
        + fila["by_only_ledger"] + fila["by_match_rounding"]
    )


def test_los_totales_cuadran_al_centavo_con_muchos_decimales():
    """Suma con Decimal: sin ella, los centavos derivan al acumular."""
    montos = ["0.07", "0.29", "1.13", "10.01", "0.11", "0.03", "99.99"]
    _, summary = build(
        [bank_row("2026-08-01", f"MOV {i}", m) for i, m in enumerate(montos)],
        [],
    )

    assert summary.iloc[0]["bank_total"] == 111.63
    assert summary.iloc[0]["unexplained"] == 0.00


def test_fila_total_solo_cuando_hay_varias_cuentas():
    _, una = build(
        [bank_row("2026-08-01", "DEP", "100.00", account="A")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1", account="A")],
    )
    assert "__total__" not in set(una["account"])

    _, varias = build(
        [bank_row("2026-08-01", "DEP", "100.00", account="A"),
         bank_row("2026-08-01", "DEP", "200.00", account="B")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1", account="A")],
    )
    assert varias.iloc[-1]["account"] == "__total__"
    assert varias.iloc[-1]["bank_total"] == 300.00


def test_total_flagged_no_cuenta_dos_veces_la_fila_total():
    _, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00", account="A"),
         bank_row("2026-08-01", "DEP", "200.00", account="B")],
        [],
    )

    assert total_flagged(summary) == 2   # una excepción por cuenta, sin duplicar


def test_las_filas_rechazadas_se_cuentan_por_cuenta():
    rejected = pd.DataFrame(
        [{"source": "bank", "row": 5, "account": "A", "date": "", "description": "X",
          "amount": "1", "reason": "reason.invalid_date"}],
        columns=REJECTED_COLUMNS,
    )
    _, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00", account="A")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1", account="A")],
        rejected=rejected,
    )

    assert summary.iloc[0]["rejected"] == 1
    assert summary.iloc[0]["flagged"] == 1


def test_resumen_vacio_no_rompe():
    _, summary = build([], [])

    assert summary.empty
    assert total_flagged(summary) == 0


# --------------------------------------------------------------------------
# Idioma
# --------------------------------------------------------------------------

@pytest.mark.parametrize("lang,esperado", [("es", "Monto banco"), ("en", "Bank amount")])
def test_localize_traduce_encabezados(lang, esperado):
    frame = pd.DataFrame([{"account": "A", "bank_amount": 1.0}])

    assert esperado in localize(frame, Translator(lang)).columns


def test_localize_traduce_la_cuenta_por_defecto():
    frame = pd.DataFrame([{"account": DEFAULT_ACCOUNT, "amount": 1.0}])

    assert localize(frame, Translator("en")).iloc[0]["Account"] == "Single account"
    assert localize(frame, Translator("es")).iloc[0]["Cuenta"] == "Cuenta única"


def test_localize_traduce_origen_y_motivo_de_las_filas_rechazadas():
    frame = pd.DataFrame([{"source": "bank", "reason": "reason.invalid_date"}])
    salida = localize(frame, Translator("en")).iloc[0]

    assert salida["Source"] == "Bank"
    assert salida["Reason"] == "Invalid or empty date"


def test_idioma_no_soportado_falla_explicitamente():
    with pytest.raises(ValueError, match="Idioma no soportado"):
        Translator("fr")


@pytest.mark.parametrize("lang,titulo", [
    ("es", "REPORTE DE RECONCILIACIÓN BANCARIA"),
    ("en", "BANK RECONCILIATION REPORT"),
])
def test_el_reporte_de_consola_respeta_el_idioma(lang, titulo):
    results, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1")],
    )
    texto = format_report(results, summary, lang=lang)

    assert titulo in texto


def test_el_reporte_de_consola_muestra_las_discrepancias():
    results, summary = build(
        [bank_row("2026-08-06", "GLOBEX", "1800.00")],
        [ledger_row("2026-08-06", "Globex", "1750.00", "INV-2")],
    )
    texto = format_report(results, summary, lang="es")

    assert "GLOBEX" in texto
    assert "INV-2" in texto
    assert "Sin explicar" in texto


def test_format_report_no_rompe_sin_datos():
    results, summary = build([], [])

    assert "RECONCILIACIÓN" in format_report(results, summary, lang="es")


# --------------------------------------------------------------------------
# Exportación a Excel
# --------------------------------------------------------------------------

def test_exporta_todas_las_pestanas(tmp_path):
    results, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1")],
    )
    destino = export_report(results, summary, tmp_path / "reporte.xlsx", lang="es")
    libro = openpyxl.load_workbook(destino)

    tr = Translator("es")
    assert libro.sheetnames == [tr.sheet(name) for name in SHEET_ORDER]


def test_las_pestanas_vacias_conservan_encabezados(tmp_path):
    """Una hoja en blanco parece un error; una con encabezados y sin filas
    comunica 'revisado, nada que reportar'."""
    results, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1")],
    )
    destino = export_report(results, summary, tmp_path / "reporte.xlsx", lang="es")
    hoja = openpyxl.load_workbook(destino)[Translator("es").sheet("discrepancies")]

    assert hoja.max_row == 1                       # solo el encabezado
    assert hoja.cell(row=1, column=1).value == "Cuenta"


def test_exporta_en_ingles(tmp_path):
    results, summary = build(
        [bank_row("2026-08-01", "DEP", "100.00")],
        [ledger_row("2026-08-01", "Dep", "100.00", "INV-1")],
    )
    destino = export_report(results, summary, tmp_path / "report.xlsx", lang="en")
    libro = openpyxl.load_workbook(destino)

    assert "Summary" in libro.sheetnames
    assert libro["Summary"].cell(row=1, column=1).value == "Account"


def test_crea_el_directorio_de_salida_si_no_existe(tmp_path):
    results, summary = build([], [])
    destino = export_report(results, summary, tmp_path / "nuevo" / "sub" / "r.xlsx")

    assert destino.exists()


def test_los_montos_llevan_formato_de_dinero(tmp_path):
    results, summary = build(
        [bank_row("2026-08-01", "DEP", "1450.00")],
        [ledger_row("2026-08-01", "Dep", "1450.00", "INV-1")],
    )
    destino = export_report(results, summary, tmp_path / "reporte.xlsx", lang="es")
    hoja = openpyxl.load_workbook(destino)[Translator("es").sheet("matches")]

    columna = [c for c in range(1, hoja.max_column + 1)
               if hoja.cell(row=1, column=c).value == "Monto"][0]
    assert hoja.cell(row=2, column=columna).number_format == "#,##0.00"
